from __future__ import annotations

import json
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

ROOT = Path.cwd()
OUT = ROOT / "paper_workspace"
TEXT_DIR = OUT / "text"

def safe_name(s: str) -> str:
    s = re.sub(r'[\\/:*?"<>|\r\n]+', '_', s)
    return s[:180]

def extract_docx(path: Path):
    from docx import Document
    doc = Document(path)
    paras = []
    full = []
    for i, p in enumerate(doc.paragraphs):
        txt = p.text
        style = p.style.name if p.style is not None else ""
        paras.append({"index": i, "style": style, "text": txt})
        full.append(txt)
    tables = []
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append({"table_index": ti, "rows": rows})
        full.append(f"\n[TABLE {ti}]\n")
        for row in rows:
            full.append("\t".join(row))
    return "\n".join(full), {"paragraphs": paras, "tables": tables}

def extract_pdf(path: Path):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    chunks = []
    for i, page in enumerate(reader.pages):
        try:
            txt = page.extract_text() or ""
        except Exception as e:
            txt = f"[EXTRACTION ERROR: {e}]"
        chunks.append(f"\n--- PAGE {i+1} ---\n{txt}")
    return "".join(chunks), {"pages": len(reader.pages)}

def main():
    zips = sorted(ROOT.glob("*.zip"))
    if not zips:
        raise SystemExit("No ZIP found in repository root")
    zip_path = zips[0]

    if OUT.exists():
        shutil.rmtree(OUT)
    TEXT_DIR.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as td:
        work = Path(td) / "archive"
        work.mkdir()
        with zipfile.ZipFile(zip_path, "r") as zf:
            zf.extractall(work)

        manifest = []
        metadata = {}
        files = [p for p in work.rglob("*") if p.is_file()]
        for idx, p in enumerate(sorted(files), start=1):
            rel = p.relative_to(work).as_posix()
            size = p.stat().st_size
            ext = p.suffix.lower()
            manifest.append({"index": idx, "path": rel, "size_bytes": size, "ext": ext})
            key = f"{idx:03d}_{safe_name(p.name)}"
            try:
                if ext == ".docx":
                    text, extra = extract_docx(p)
                    (TEXT_DIR / f"{key}.txt").write_text(text[:1500000], encoding="utf-8", errors="replace")
                    (TEXT_DIR / f"{key}.structure.json").write_text(
                        json.dumps(extra, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
                    metadata[rel] = {"extract": f"text/{key}.txt", "structure": f"text/{key}.structure.json", **extra}
                elif ext == ".pdf":
                    text, extra = extract_pdf(p)
                    (TEXT_DIR / f"{key}.txt").write_text(text[:1000000], encoding="utf-8", errors="replace")
                    metadata[rel] = {"extract": f"text/{key}.txt", **extra}
                elif ext in {".txt", ".md", ".csv", ".rtf"}:
                    raw = p.read_text(encoding="utf-8", errors="replace")
                    (TEXT_DIR / f"{key}.txt").write_text(raw[:1000000], encoding="utf-8")
                    metadata[rel] = {"extract": f"text/{key}.txt"}
            except Exception as e:
                metadata[rel] = {"error": repr(e)}

    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["index\tsize_bytes\text\tpath"]
    for row in manifest:
        lines.append(f"{row['index']}\t{row['size_bytes']}\t{row['ext']}\t{row['path']}")
    (OUT / "manifest.tsv").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print((OUT / "manifest.tsv").read_text(encoding="utf-8"))

if __name__ == "__main__":
    main()
